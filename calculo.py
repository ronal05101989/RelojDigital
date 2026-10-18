"""Cálculo de asistencia: tardanzas, faltas, permisos, horas trabajadas y horas extra."""
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO

DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
MARCAJE_DOBLE_MIN = 5      # marcajes con menos de 5 min de diferencia cuentan como uno solo
DESCUENTA_DESCANSO_H = 6   # el descanso solo se descuenta si estuvo 6 h o más

TIPOS_AUSENCIA = ["Permiso personal", "Permiso por salud", "Licencia médica", "Vacaciones",
                  "Comisión de servicio", "Capacitación", "Licencia por maternidad",
                  "Licencia por paternidad", "Licencia por fallecimiento", "Licencia sin goce", "Otro"]

REPORTES = {
    "resumen": "Resumen por empleado",
    "entradas": "Entradas y salidas",
    "inasistencias": "Inasistencias",
    "tardanzas": "Tardanzas",
    "extras": "Horas extra",
    "permisos": "Permisos y licencias",
}


def a_minutos(hhmm):
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def hm(minutos):
    """125 -> '2:05'"""
    if not minutos:
        return ""
    return f"{int(minutos) // 60}:{int(minutos) % 60:02d}"


def detalle(con, desde, hasta, codigo=None, departamento=None):
    """Una fila por empleado y día con su estado."""
    sql = ("SELECT e.codigo, e.nombre, e.departamento, h.nombre AS horario, h.entrada AS h_entrada, "
           "h.salida AS h_salida, h.tolerancia_min, h.descanso_min, h.dias "
           "FROM empleados e LEFT JOIN horarios h ON h.id = e.horario_id WHERE e.activo = 1")
    params = []
    if codigo:
        sql += " AND e.codigo = ?"
        params.append(codigo)
    if departamento:
        sql += " AND e.departamento = ?"
        params.append(departamento)
    empleados = con.execute(sql + " ORDER BY e.nombre", params).fetchall()

    marcas = defaultdict(list)
    for m in con.execute("SELECT codigo, fecha_hora FROM marcajes WHERE fecha_hora >= ? AND fecha_hora < ? "
                         "ORDER BY fecha_hora", (desde.isoformat(), (hasta + timedelta(days=1)).isoformat())):
        momento = datetime.strptime(m["fecha_hora"], "%Y-%m-%d %H:%M:%S")
        marcas[(m["codigo"], momento.date())].append(momento)

    feriados = {r["fecha"]: r["descripcion"] for r in con.execute(
        "SELECT * FROM feriados WHERE fecha BETWEEN ? AND ?", (desde.isoformat(), hasta.isoformat()))}
    ausencias = defaultdict(list)
    for a in con.execute("SELECT * FROM ausencias WHERE desde <= ? AND hasta >= ? ORDER BY hora_desde",
                         (hasta.isoformat(), desde.isoformat())):
        ausencias[a["codigo"]].append(a)

    fila_min = con.execute("SELECT valor FROM ajustes WHERE clave = 'extra_minimo_min'").fetchone()
    extra_minimo = int(fila_min["valor"]) if fila_min and fila_min["valor"].isdigit() else 30

    hoy, ahora = date.today(), datetime.now()
    filas = []
    for e in empleados:
        dias_lab = {int(d) for d in (e["dias"] or "").split(",") if d.strip()}
        d = desde
        while d <= hasta:
            ms = marcas.get((e["codigo"], d), [])
            iso = d.isoformat()
            aus = next((a for a in ausencias[e["codigo"]] if a["desde"] <= iso <= a["hasta"]), None)
            dia_completo = aus is not None and not aus["hora_desde"]
            laborable = e["h_entrada"] is not None and d.weekday() in dias_lab and iso not in feriados

            # Permiso por horas al inicio de la jornada: se le espera al terminar el permiso
            h_entrada = e["h_entrada"] or ""
            if laborable and aus and not dia_completo and aus["hora_desde"] <= h_entrada < aus["hora_hasta"]:
                h_entrada = aus["hora_hasta"]

            fila = {"codigo": e["codigo"], "nombre": e["nombre"], "departamento": e["departamento"],
                    "horario": e["horario"] or "Sin horario", "fecha": d, "dia": DIAS[d.weekday()],
                    "h_entrada": h_entrada if laborable else "", "entrada": "", "salida": "",
                    "marcajes": len(ms), "tardanza": 0, "trabajado": 0, "extra": 0, "estado": "",
                    "ausencia": "", "detalle_ausencia": "", "es_permiso": False}
            if aus:
                fila["ausencia"] = aus["tipo"] if dia_completo else \
                    f"{aus['tipo']} ({aus['hora_desde']}–{aus['hora_hasta']})"
                fila["detalle_ausencia"] = " · ".join(x for x in (aus["documento"], aus["motivo"]) if x)

            if ms:
                ent, sal = ms[0], ms[-1]
                completo = (sal - ent).total_seconds() >= MARCAJE_DOBLE_MIN * 60
                fila["entrada"] = ent.strftime("%H:%M")
                if laborable and not dia_completo:
                    retraso = ent.hour * 60 + ent.minute - a_minutos(h_entrada)
                    if retraso > e["tolerancia_min"]:
                        fila["tardanza"] = retraso
                if completo:
                    fila["salida"] = sal.strftime("%H:%M")
                    bruto = (sal - ent).total_seconds() // 60
                    descanso = (e["descanso_min"] or 0) if bruto >= DESCUENTA_DESCANSO_H * 60 else 0
                    fila["trabajado"] = max(0, bruto - descanso)
                    if laborable:
                        jornada = a_minutos(e["h_salida"]) - a_minutos(e["h_entrada"]) - (e["descanso_min"] or 0)
                        fila["extra"] = max(0, fila["trabajado"] - jornada)
                        fila["estado"] = "Tardanza" if fila["tardanza"] else "Puntual"
                    else:
                        fila["extra"] = fila["trabajado"]
                        fila["estado"] = "Feriado trabajado" if iso in feriados else "Día libre trabajado"
                    if fila["extra"] < extra_minimo:  # unos minutos de más no son horas extra
                        fila["extra"] = 0
                elif d == hoy:
                    fila["estado"] = "En turno" if not fila["tardanza"] else "En turno (tarde)"
                else:
                    fila["estado"] = "Sin marcaje de salida"
            elif laborable and dia_completo:
                fila["estado"] = aus["tipo"]
                fila["es_permiso"] = True
            elif laborable:
                inicio = datetime.combine(d, datetime.strptime(h_entrada, "%H:%M").time())
                if inicio + timedelta(minutes=e["tolerancia_min"]) < ahora:
                    fila["estado"] = "Falta"
                elif d >= hoy:
                    fila["estado"] = "Pendiente"
            # "Pendiente" (aún no es su hora de entrada) solo se muestra para hoy
            if fila["estado"] and (fila["estado"] != "Pendiente" or d == hoy):
                filas.append(fila)
            d += timedelta(days=1)
    return filas


def resumen(filas):
    """Totales por empleado a partir del detalle."""
    tot = {}
    for f in filas:
        r = tot.setdefault(f["codigo"], {"codigo": f["codigo"], "nombre": f["nombre"],
                                         "departamento": f["departamento"], "dias_trabajados": 0,
                                         "faltas": 0, "tardanzas": 0, "min_tardanza": 0, "permisos": 0,
                                         "sin_salida": 0, "trabajado": 0, "extra": 0})
        if f["marcajes"]:
            r["dias_trabajados"] += 1
        if f["estado"] == "Falta":
            r["faltas"] += 1
        if f["estado"] == "Sin marcaje de salida":
            r["sin_salida"] += 1
        if f["es_permiso"]:
            r["permisos"] += 1
        if f["tardanza"]:
            r["tardanzas"] += 1
            r["min_tardanza"] += f["tardanza"]
        r["trabajado"] += f["trabajado"]
        r["extra"] += f["extra"]
    return sorted(tot.values(), key=lambda r: r["nombre"])


def filtrar(filas, tipo):
    if tipo == "entradas":
        return [f for f in filas if f["marcajes"]]
    if tipo == "inasistencias":
        return [f for f in filas if f["estado"] == "Falta"]
    if tipo == "tardanzas":
        return [f for f in filas if f["tardanza"]]
    if tipo == "extras":
        return [f for f in filas if f["extra"]]
    if tipo == "permisos":
        return [f for f in filas if f["ausencia"]]
    return filas


# Columnas de cada reporte: (título, función que saca el valor, es_numérica)
_BASE = [("Fecha", lambda f: f["fecha"], False), ("Día", lambda f: f["dia"], False),
         ("Código", lambda f: f["codigo"], False), ("Nombre", lambda f: f["nombre"], False),
         ("Departamento", lambda f: f["departamento"], False)]
COLUMNAS = {
    "entradas": _BASE + [("Horario", lambda f: f["horario"], False), ("Entrada", lambda f: f["entrada"], False),
                         ("Salida", lambda f: f["salida"], False), ("Trabajado", lambda f: hm(f["trabajado"]), True),
                         ("Extra", lambda f: hm(f["extra"]), True), ("Estado", lambda f: f["estado"], False),
                         ("Permiso", lambda f: f["ausencia"], False)],
    "inasistencias": _BASE + [("Horario", lambda f: f["horario"], False),
                              ("Debía entrar", lambda f: f["h_entrada"], False),
                              ("Permiso parcial", lambda f: f["ausencia"], False)],
    "tardanzas": _BASE + [("Debía entrar", lambda f: f["h_entrada"], False),
                          ("Marcó", lambda f: f["entrada"], False),
                          ("Minutos tarde", lambda f: f["tardanza"], True),
                          ("Permiso", lambda f: f["ausencia"], False)],
    "extras": _BASE + [("Entrada", lambda f: f["entrada"], False), ("Salida", lambda f: f["salida"], False),
                       ("Trabajado", lambda f: hm(f["trabajado"]), True),
                       ("Horas extra", lambda f: hm(f["extra"]), True), ("Tipo de día", lambda f: f["estado"], False)],
    "permisos": _BASE + [("Permiso / licencia", lambda f: f["ausencia"], False),
                         ("Documento / motivo", lambda f: f["detalle_ausencia"], False),
                         ("Estado del día", lambda f: f["estado"], False)],
    "resumen": [("Código", lambda r: r["codigo"], False), ("Nombre", lambda r: r["nombre"], False),
                ("Departamento", lambda r: r["departamento"], False),
                ("Días trabajados", lambda r: r["dias_trabajados"], True), ("Faltas", lambda r: r["faltas"], True),
                ("Tardanzas", lambda r: r["tardanzas"], True),
                ("Tiempo tarde", lambda r: hm(r["min_tardanza"]), True),
                ("Días de permiso", lambda r: r["permisos"], True), ("Sin salida", lambda r: r["sin_salida"], True),
                ("Horas trabajadas", lambda r: hm(r["trabajado"]), True),
                ("Horas extra", lambda r: hm(r["extra"]), True)],
}


def tabla(filas, tipo):
    """Devuelve (columnas, datos, total_texto) del reporte pedido."""
    datos = resumen(filas) if tipo == "resumen" else filtrar(filas, tipo)
    total = f"{len(datos)} registros"
    if tipo == "tardanzas":
        total += f" · {hm(sum(f['tardanza'] for f in datos)) or '0:00'} h de tardanza en total"
    elif tipo == "extras":
        total += f" · {hm(sum(f['extra'] for f in datos)) or '0:00'} horas extra en total"
    elif tipo == "resumen":
        total = f"{len(datos)} empleados"
    return COLUMNAS[tipo], datos, total


def excel(filas, desde, hasta, tipos, subtitulo=""):
    """Genera un Excel con una hoja por cada tipo de reporte. Devuelve bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E78")
    colores = {"Falta": PatternFill("solid", fgColor="F8D7DA"),
               "Tardanza": PatternFill("solid", fgColor="FFF3CD"),
               "Sin marcaje de salida": PatternFill("solid", fgColor="FFF3CD")}

    wb = Workbook()
    wb.remove(wb.active)
    for tipo in tipos:
        columnas, datos, total = tabla(filas, tipo)
        ws = wb.create_sheet(REPORTES[tipo][:31])
        ws.append([f"{REPORTES[tipo]} del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([" · ".join(x for x in (subtitulo, total) if x)])
        ws.append([c[0] for c in columnas])
        for celda in ws[3]:
            celda.font, celda.fill = cabecera, relleno
            celda.alignment = Alignment(horizontal="center")
        for d in datos:
            ws.append([c[1](d) if c[1](d) != "" else None for c in columnas])
            color = colores.get(d.get("estado")) if tipo == "entradas" else None
            if color:
                for celda in ws[ws.max_row]:
                    celda.fill = color
        for i, (titulo, _, _) in enumerate(columnas, start=1):
            letra = get_column_letter(i)
            largo = max([len(titulo)] + [len(str(c.value or "")) for c in ws[letra][3:]])
            ws.column_dimensions[letra].width = min(45, largo + 3)
            if titulo == "Fecha":
                for c in ws[letra][3:]:
                    c.number_format = "DD/MM/YYYY"
                ws.column_dimensions[letra].width = 12
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(columnas))}{max(ws.max_row, 3)}"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
